#!/usr/bin/env python3
# Copyright 2026 The VeriHarness Authors.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Recalculate a workbook after editing it with openpyxl, and report what changed.

    python3 xlsx_recalc.py EDITED.xlsx [--baseline BEFORE.xlsx] [--inplace]

openpyxl writes formulas without cached values; graders and judges read cached
values. This
runs headless LibreOffice on a copy (convert-to xlsx), which recomputes every
formula, then
prints per sheet: cells whose cached value changed versus --baseline (the file
before your
edit), so you can confirm that only the cells you meant to touch — and their
dependents —
moved. With --inplace the recalculated file replaces EDITED.xlsx (a
.pre-recalc.xlsx copy is
kept beside it). Charts and some formatting may not survive a LibreOffice round
trip: run
xlsx_diff.py on the result before delivering, and keep the openpyxl version if
it did damage.
"""

import argparse, os, shutil, subprocess, sys, tempfile
from pathlib import Path
import openpyxl

ap = argparse.ArgumentParser(
    description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
)
ap.add_argument("edited")
ap.add_argument("--baseline")
ap.add_argument("--inplace", action="store_true")
ap.add_argument("--limit", type=int, default=60)
a = ap.parse_args()
src = Path(a.edited).resolve()
with tempfile.TemporaryDirectory(prefix="xlsx_recalc_", dir="/tmp") as td:
  work = Path(td) / src.name
  shutil.copyfile(src, work)
  env = dict(os.environ, HOME=td)
  r = subprocess.run(
      [
          "soffice",
          "--headless",
          "--calc",
          "--convert-to",
          "xlsx",
          "--outdir",
          str(Path(td) / "out"),
          str(work),
      ],
      capture_output=True,
      text=True,
      timeout=300,
      env=env,
  )
  out = Path(td) / "out" / src.name
  if r.returncode != 0 or not out.exists():
    print("recalc FAILED:", (r.stderr or r.stdout)[-400:])
    sys.exit(1)
  recalced = out.read_bytes()
  base_path = Path(a.baseline) if a.baseline else None
  wb_new = openpyxl.load_workbook(out, data_only=True)
  wb_old = (
      openpyxl.load_workbook(base_path, data_only=True) if base_path else None
  )
  wb_fml = openpyxl.load_workbook(out, data_only=False)
  changed = 0
  uncached = 0
  for ws in wb_new.worksheets:
    wf = wb_fml[ws.title]
    wo = wb_old[ws.title] if wb_old and ws.title in wb_old.sheetnames else None
    for row in ws.iter_rows():
      for c in row:
        f = wf[c.coordinate].value
        if isinstance(f, str) and f.startswith("=") and c.value is None:
          uncached += 1
        if wo is None:
          continue
        o = wo[c.coordinate].value
        if o != c.value and not (
            isinstance(o, (int, float))
            and isinstance(c.value, (int, float))
            and abs(o - c.value) <= 1e-9 * max(1, abs(o))
        ):
          changed += 1
          if changed <= a.limit:
            print(
                f"{ws.title}!{c.coordinate}: {o!r} -> {c.value!r}  "
                f" {f if isinstance(f, str) and f.startswith('=') else ''}"
            )
  if wb_old is not None:
    print(
        f"# cells whose cached value changed vs baseline: {changed}"
        + (f" (showing {a.limit})" if changed > a.limit else "")
    )
  print(f"# formulas still without a cached value after recalc: {uncached}")
  if a.inplace:
    shutil.copyfile(src, src.with_name(src.stem + ".pre-recalc.xlsx"))
    src.write_bytes(recalced)
    print(
        f"# replaced {src.name} with the recalculated file (previous version"
        f" kept as {src.stem}.pre-recalc.xlsx)"
    )
  else:
    dst = src.with_name(src.stem + ".recalc.xlsx")
    dst.write_bytes(recalced)
    print(f"# recalculated copy written to {dst}")
