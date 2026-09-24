#!/usr/bin/env python3
# Copyright 2026 The VeriHarness Authors.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Formulas and cached values side by side.

usage: xlsx_dump.py FILE [SHEET] [RANGE] [--all]      e.g. xlsx_dump.py
model.xlsx LBO A1:H40

Prints the workbook's calculation settings, then one line per non-empty cell:
COORD <TAB> formula <TAB> => cached value   (or COORD <TAB> literal value).
Without RANGE, each sheet is cut after its first 200 rows (pass --all to lift).
"""

import signal
import sys

ROW_CAP = 200

import openpyxl

signal.signal(signal.SIGPIPE, signal.SIG_DFL)

if len(sys.argv) < 2:
  sys.exit(__doc__)
show_all = "--all" in sys.argv
argv = [a for a in sys.argv if a != "--all"]
path = argv[1]
sheet = argv[2] if len(argv) > 2 else None
rng = argv[3] if len(argv) > 3 else None

wbf = openpyxl.load_workbook(path, data_only=False)
wbv = openpyxl.load_workbook(path, data_only=True)
calc = wbf.calculation
print(f"# {path}")
print(f"# sheets: {wbf.sheetnames}")
if calc is not None:
  print(
      f"# calcPr: iterate={calc.iterate} iterateCount={calc.iterateCount}"
      f" iterateDelta={calc.iterateDelta} fullCalcOnLoad={calc.fullCalcOnLoad}"
      f" calcMode={calc.calcMode}"
  )
else:
  print("# calcPr: None")
for sn in ([sheet] if sheet else wbf.sheetnames):
  wsf, wsv = wbf[sn], wbv[sn]
  cells = wsf[rng] if rng else wsf.iter_rows()
  if rng and ":" not in rng:
    cells = ((cells,),)
  print(f"\n## {sn}  (dims {wsf.dimensions})")
  n = 0
  for r_i, row in enumerate(cells, 1):
    if not rng and not show_all and r_i > ROW_CAP:
      print(
          f"# ... cut at row {ROW_CAP} of {wsf.max_row}; pass a RANGE or --all"
          " for the rest"
      )
      break
    for c in row:
      if c.value is None:
        continue
      n += 1
      if isinstance(c.value, str) and c.value.startswith("="):
        print(f"{c.coordinate}\t{c.value}\t=> {wsv[c.coordinate].value!r}")
      else:
        print(f"{c.coordinate}\t{c.value!r}")
  print(f"# {n} non-empty cells")
