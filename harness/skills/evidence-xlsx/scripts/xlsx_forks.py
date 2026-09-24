# Copyright 2026 The VeriHarness Authors.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

#!/usr/bin/env python3
# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     https://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""xlsx_forks.py FILE1.xlsx FILE2.xlsx [FILE3.xlsx ...]

Where the candidates store different things in the same cell: a locator for
disagreements,
no verdict. For every cell that differs across the files, prints the camps
(formula or value,
who holds each) and flags the kind of split when it is recognisable:

  SIGN    one camp stores a leading minus or a reversed subtraction, the other
  stores it plain
  BLANK   some candidates fill the cell, others leave it empty
  FORMULA the formulas differ in shape (different operands or operators)
  VALUE   same formula shape, different references or constants

Compare the camps against the input workbook's own notes, labels and filled
sibling rows;
this script only shows where the split is.
"""

from collections import defaultdict
import re
import sys

import openpyxl

MINUS = re.compile(
    r"^=\s*-|^=\s*\(?\s*[A-Z$0-9!']+\s*-\s*[A-Z$0-9!']+\s*\)?\s*$"
)


def shape(v):
  return re.sub(r"\$?[A-Z]{1,3}\$?\d+|\d+(?:\.\d+)?", "#", str(v).upper())


def label(ws, r):
  for c in range(1, min(ws.max_column, 4) + 1):
    v = ws.cell(r, c).value
    if isinstance(v, str) and v.strip() and not str(v).startswith("="):
      return v.strip()[:40]
  return ""


def main(paths):
  books = {p: openpyxl.load_workbook(p, data_only=False) for p in paths}
  cells = defaultdict(dict)  # (sheet, coord) -> {path: value}
  for p, wb in books.items():
    for ws in wb.worksheets:
      for row in ws.iter_rows():
        for c in row:
          if c.value is not None:
            cells[(ws.title, c.coordinate)][p] = c.value
  n = 0
  for (sheet, coord), by in sorted(cells.items()):
    vals = {p: by.get(p) for p in paths}
    if len({str(v) for v in vals.values()}) < 2:
      continue
    camps = defaultdict(list)
    for p, v in vals.items():
      camps[str(v) if v is not None else ""].append(p)
    kinds = set()
    if "" in camps:
      kinds.add("BLANK")
    forms = [k for k in camps if k.startswith("=")]
    if len(forms) >= 2:
      signs = {bool(MINUS.match(f)) for f in forms}
      kinds.add(
          "SIGN"
          if len(signs) == 2
          else ("FORMULA" if len({shape(f) for f in forms}) > 1 else "VALUE")
      )
    elif len(camps) - ("" in camps) >= 2:
      kinds.add("VALUE")
    ws = (
        next(iter(books.values()))[sheet]
        if sheet in next(iter(books.values())).sheetnames
        else None
    )
    lbl = label(ws, int(re.sub(r"[A-Z]", "", coord))) if ws is not None else ""
    print(f"[{'/'.join(sorted(kinds)) or 'DIFF'}] {sheet}!{coord} [{lbl}]")
    for k, members in sorted(camps.items(), key=lambda kv: -len(kv[1])):
      names = " ".join(
          m.split("/")[-3] if "/deliverables/" in m else m for m in members
      )
      print(f"    {len(members):2d}  {k[:90] if k else '(blank)'}   <- {names}")
    n += 1
    if n >= 400:
      print("... (cut at 400 cells)")
      break
  if not n:
    print("no differing cells")
  return 0


if __name__ == "__main__":
  if len(sys.argv) < 3:
    print(__doc__)
    sys.exit(1)
  sys.exit(main(sys.argv[1:]))
