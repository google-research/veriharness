#!/usr/bin/env python3
# Copyright 2026 The VeriHarness Authors.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Cell-by-cell diff of two workbooks: formulas, cached values AND formatting.

usage: xlsx_diff.py A.xlsx B.xlsx [--limit N] [--no-format]

Reports sheets present in only one file, then for each shared sheet every cell
where the formula or the cached value differs (numeric tolerance 1e-9 relative),
followed by a separate FORMAT section listing cells whose font colour, bold,
fill or number format differ (graders check colour coding and number formats;
a content diff alone is blind to them). --no-format skips that section.
Output per cell:  SHEET!A1 | A: <formula> => value | B: <formula> => value
"""

import math
import signal
import sys

import openpyxl

signal.signal(signal.SIGPIPE, signal.SIG_DFL)

argv = sys.argv[1:]
limit = 300
if "--limit" in argv:
  i = argv.index("--limit")
  limit = int(argv[i + 1])
  del argv[i : i + 2]
do_format = "--no-format" not in argv
args = [x for x in argv if not x.startswith("--")]
if len(args) != 2:
  sys.exit(__doc__)
a, b = args


def load(p):
  return openpyxl.load_workbook(p, data_only=False), openpyxl.load_workbook(
      p, data_only=True
  )


def same(x, y):
  if (
      isinstance(x, (int, float))
      and isinstance(y, (int, float))
      and not isinstance(x, bool)
      and not isinstance(y, bool)
  ):
    return math.isclose(x, y, rel_tol=1e-9, abs_tol=1e-12)
  return x == y


def fmt(c):
  f = c.font
  col = (
      f.color.rgb
      if f is not None and f.color is not None and isinstance(f.color.rgb, str)
      else None
  )
  fill = (
      c.fill.fgColor.rgb
      if c.fill is not None
      and c.fill.fill_type
      and isinstance(c.fill.fgColor.rgb, str)
      else None
  )
  return (col, bool(f.bold) if f is not None else False, fill, c.number_format)


def cell(wsf, wsv, coord):
  f = wsf[coord].value
  v = wsv[coord].value
  if isinstance(f, str) and f.startswith("="):
    return f, v, f"{f} => {v!r}"
  return None, f, repr(f)


af, av = load(a)
bf, bv = load(b)
only_a = [s for s in af.sheetnames if s not in bf.sheetnames]
only_b = [s for s in bf.sheetnames if s not in af.sheetnames]
if only_a:
  print(f"# sheets only in A: {only_a}")
if only_b:
  print(f"# sheets only in B: {only_b}")
total = 0
for sn in [s for s in af.sheetnames if s in bf.sheetnames]:
  wsaf, wsav, wsbf, wsbv = af[sn], av[sn], bf[sn], bv[sn]
  coords = set()
  for ws in (wsaf, wsbf):
    for row in ws.iter_rows():
      for c in row:
        if c.value is not None:
          coords.add(c.coordinate)
  diffs = []
  for coord in sorted(coords, key=lambda k: (wsaf[k].row, wsaf[k].column)):
    fa, va, sa = cell(wsaf, wsav, coord)
    fb, vb, sb = cell(wsbf, wsbv, coord)
    if fa != fb or not same(va, vb):
      diffs.append(f"{sn}!{coord} | A: {sa} | B: {sb}")
  total += len(diffs)
  print(f"\n## {sn}: {len(diffs)} differing cells")
  for line in diffs[:limit]:
    print(line)
  if len(diffs) > limit:
    print(f"... {len(diffs) - limit} more (raise --limit)")
  if do_format:
    fdiffs = []
    for coord in sorted(coords, key=lambda k: (wsaf[k].row, wsaf[k].column)):
      fa, fb = fmt(wsaf[coord]), fmt(wsbf[coord])
      if fa != fb:
        names = ("font", "bold", "fill", "numfmt")
        changed = ", ".join(
            f"{n}: {x!r} -> {y!r}" for n, x, y in zip(names, fa, fb) if x != y
        )
        fdiffs.append(f"{sn}!{coord} | FORMAT {changed}")
    print(f"## {sn}: {len(fdiffs)} cells with formatting differences")
    for line in fdiffs[:limit]:
      print(line)
    if len(fdiffs) > limit:
      print(f"... {len(fdiffs) - limit} more (raise --limit)")
print(f"\n# total differing cells (content): {total}")
