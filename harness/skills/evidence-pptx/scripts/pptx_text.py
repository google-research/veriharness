#!/usr/bin/env python3
# Copyright 2026 The VeriHarness Authors.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Slide-by-slide contents of a .pptx: titles, text frames (in shape order), tables cell by

cell, speaker notes — and an explicit line for every shape that could not be
read.

    python3 pptx_text.py FILE [--slides 3,5-7] [--no-notes]

Text inside tables, grouped shapes and placeholders is easy to miss or to read
out of
order with ad-hoc code; this walks every shape (recursing into groups), prints
what it
holds, and marks pictures/charts/unreadable shapes as such instead of skipping
them.
"""

import argparse, signal, sys

signal.signal(signal.SIGPIPE, signal.SIG_DFL)
from pptx import Presentation
from pptx.util import Emu

ap = argparse.ArgumentParser(
    description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
)
ap.add_argument("file")
ap.add_argument("--slides", help="e.g. 3,5-7 (1-based)")
ap.add_argument("--no-notes", action="store_true")
a = ap.parse_args()


def want(i):
  if not a.slides:
    return True
  for part in a.slides.split(","):
    if "-" in part:
      lo, hi = part.split("-")
      if int(lo) <= i <= int(hi):
        return True
    elif int(part) == i:
      return True
  return False


def walk(shapes, depth=0):
  pad = "  " * depth
  for sh in shapes:
    kind = sh.shape_type
    name = sh.name
    try:
      if sh.shape_type is not None and str(kind).startswith("GROUP"):
        print(f"{pad}[group {name}]")
        walk(sh.shapes, depth + 1)
        continue
      if getattr(sh, "has_table", False) and sh.has_table:
        t = sh.table
        print(
            f"{pad}[table {name}: {len(t.rows)} rows x {len(t.columns)} cols]"
        )
        for r in t.rows:
          print(
              pad
              + "  "
              + "\t".join(c.text.replace("\n", " ").strip() for c in r.cells)
          )
        continue
      if getattr(sh, "has_chart", False) and sh.has_chart:
        ch = sh.chart
        try:
          cats = list(ch.plots[0].categories) if ch.plots else []
          print(
              f"{pad}[chart {name}: type={ch.chart_type}"
              f" categories={cats[:12]}]"
          )
          for pl in ch.plots:
            for se in pl.series:
              print(f"{pad}  series {se.name}: {list(se.values)[:12]}")
        except Exception as e:  # noqa: BLE001
          print(
              f"{pad}[chart {name}: could not read series ({type(e).__name__})"
              " — render the slide]"
          )
        continue
      if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
        txt = "\n".join(p.text for p in sh.text_frame.paragraphs).strip()
        ph = " placeholder" if sh.is_placeholder else ""
        if txt:
          print(f"{pad}[text {name}{ph}]")
          [print(pad + "  " + l) for l in txt.splitlines()]
        else:
          print(f"{pad}[text {name}{ph}: empty]")
        continue
      if str(kind).startswith("PICTURE"):
        print(
            f"{pad}[picture {name}:"
            f" {Emu(sh.width).inches:.1f}x{Emu(sh.height).inches:.1f} in —"
            " content not readable as text; render the slide]"
        )
        continue
      print(f"{pad}[shape {name}: type={kind}, no text]")
    except Exception as e:  # noqa: BLE001
      print(
          f"{pad}[shape {name}: UNREADABLE ({type(e).__name__}: {e}) — render"
          " the slide to see it]"
      )


try:
  prs = Presentation(a.file)
except Exception as e:  # noqa: BLE001
  print(f"cannot open {a.file}: {type(e).__name__}: {e}")
  sys.exit(1)
for i, slide in enumerate(prs.slides, 1):
  if not want(i):
    continue
  title = (
      slide.shapes.title.text.strip()
      if slide.shapes.title is not None and slide.shapes.title.has_text_frame
      else ""
  )
  print(
      f"=== slide {i} (layout:"
      f" {slide.slide_layout.name}){' — ' + title if title else ''}"
  )
  walk(slide.shapes)
  if not a.no_notes and slide.has_notes_slide:
    n = (
        slide.notes_slide.notes_text_frame.text.strip()
        if slide.notes_slide.notes_text_frame
        else ""
    )
    if n:
      print("[notes]")
      [print("  " + l) for l in n.splitlines()]
  print()
print(f"# {len(prs.slides)} slides total")
